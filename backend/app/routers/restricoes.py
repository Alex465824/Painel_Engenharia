"""Restricoes de obra.

A API do Prevision nao publica restricoes (nenhum dos 16 endpoints da especificacao), embora o
produto as tenha. Para o painel ter a aba funcionando, ficam registradas aqui por cadastro ou
por importacao de planilha exportada do Prevision. Lembrete de risco: isso cria um segundo
lugar onde a restricao existe - o ideal e importar do Prevision, nao digitar duas vezes.
"""

from __future__ import annotations

import csv
import io
import unicodedata
from collections import Counter
from datetime import date, datetime

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.database import get_db
from app.db_models import ObraCache, Restricao

router = APIRouter(prefix="/api/restricoes", tags=["restricoes"])

STATUS = ("Nao iniciada", "Iniciada", "Concluida")
CAMPOS = ("descricao", "categoria", "responsavel", "prazo", "status", "criada_em", "concluida_em", "observacao")
SINONIMOS = {
    "descricao": ["descri", "restri", "titulo", "assunto", "pendencia", "impedimento"],
    "categoria": ["categoria", "tipo", "classifica", "natureza"],
    "responsavel": ["respons", "dono", "owner", "atribu"],
    "prazo": ["prazo", "limite", "previst", "vencimento"],
    "status": ["status", "situa"],
    "concluida_em": ["conclus", "concluid", "finaliz", "resolv"],
    "criada_em": ["criad", "abertura", "identific", "inclus"],
    "observacao": ["observ", "coment", "detalhe", "nota"],
}


def _sem_acento(t: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", str(t or "")) if unicodedata.category(c) != "Mn").lower().strip()


def _data(valor) -> date | None:
    if valor in (None, ""):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor).strip()[:10]
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    return None


def _status(valor) -> str:
    t = _sem_acento(valor)
    if "nao" in t and "inici" in t:
        return "Nao iniciada"
    if any(k in t for k in ("conclu", "finaliz", "resolv", "encerr", "feito")):
        return "Concluida"
    if any(k in t for k in ("inici", "andamento", "tratativa", "curso")):
        return "Iniciada"
    return "Nao iniciada"


class RestricaoIn(BaseModel):
    obra_id: int
    descricao: str
    categoria: str | None = None
    responsavel: str | None = None
    prazo: str | None = None
    status: str = "Nao iniciada"
    criada_em: str | None = None
    concluida_em: str | None = None
    observacao: str | None = None


def _situacao(r: Restricao, hoje: date) -> str:
    if r.status == "Concluida":
        return "concluida"
    if r.prazo is None:
        return "sem_prazo"
    return "em_atraso" if r.prazo < hoje else "no_prazo"


def _saida(r: Restricao, hoje: date) -> dict:
    return {
        "id": r.id, "obra_id": r.obra_id, "descricao": r.descricao, "categoria": r.categoria,
        "responsavel": r.responsavel, "prazo": r.prazo.isoformat() if r.prazo else None,
        "status": r.status, "criada_em": r.criada_em.isoformat() if r.criada_em else None,
        "concluida_em": r.concluida_em.isoformat() if r.concluida_em else None,
        "origem": r.origem, "observacao": r.observacao, "situacao": _situacao(r, hoje),
        "dias_atraso": (hoje - r.prazo).days if r.prazo and r.status != "Concluida" and r.prazo < hoje else None,
    }


def _aplicar(r: Restricao, c: RestricaoIn) -> None:
    if not c.descricao.strip():
        raise HTTPException(400, "Descreva a restricao.")
    if c.status not in STATUS:
        raise HTTPException(400, f"Status invalido. Use: {', '.join(STATUS)}")
    r.obra_id, r.descricao = c.obra_id, c.descricao.strip()
    r.categoria = (c.categoria or "").strip() or None
    r.responsavel = (c.responsavel or "").strip() or None
    r.prazo = _data(c.prazo)
    r.status = c.status
    r.criada_em = _data(c.criada_em) or r.criada_em or date.today()
    r.concluida_em = _data(c.concluida_em) or (date.today() if c.status == "Concluida" and not r.concluida_em else r.concluida_em)
    if c.status != "Concluida":
        r.concluida_em = None
    r.observacao = (c.observacao or "").strip() or None


@router.get("")
def listar(obra_id: int | None = None, db: Session = Depends(get_db)):
    consulta = select(Restricao).order_by(Restricao.prazo.is_(None), Restricao.prazo.asc(), Restricao.id.desc())
    if obra_id:
        consulta = consulta.where(Restricao.obra_id == obra_id)
    hoje = date.today()
    return [_saida(r, hoje) for r in db.scalars(consulta)]


@router.get("/resumo")
def resumo(obra_id: int | None = None, db: Session = Depends(get_db)):
    consulta = select(Restricao)
    if obra_id:
        consulta = consulta.where(Restricao.obra_id == obra_id)
    itens = list(db.scalars(consulta))
    hoje = date.today()
    abertas = [r for r in itens if r.status != "Concluida"]
    concluidas = [r for r in itens if r.status == "Concluida"]
    com_data = [r for r in concluidas if r.prazo and r.concluida_em]
    conc_atraso = [r for r in com_data if r.concluida_em > r.prazo]
    return {
        "total": len(itens),
        "nao_iniciadas": sum(1 for r in itens if r.status == "Nao iniciada"),
        "iniciadas": sum(1 for r in itens if r.status == "Iniciada"),
        "concluidas": len(concluidas),
        "abertas_em_atraso": sum(1 for r in abertas if r.prazo and r.prazo < hoje),
        "abertas_no_prazo": sum(1 for r in abertas if not r.prazo or r.prazo >= hoje),
        "concluidas_em_atraso": len(conc_atraso),
        "concluidas_no_prazo": len(com_data) - len(conc_atraso),
        "concluidas_sem_data": len(concluidas) - len(com_data),
        "por_categoria_abertas": [
            {"categoria": c, "n": n}
            for c, n in Counter((r.categoria or "Sem categoria") for r in abertas).most_common()
        ],
    }


@router.post("", status_code=201)
def criar(corpo: RestricaoIn, db: Session = Depends(get_db)):
    if db.get(ObraCache, corpo.obra_id) is None:
        raise HTTPException(404, "Obra nao encontrada")
    r = Restricao()
    _aplicar(r, corpo)
    db.add(r)
    db.commit()
    return _saida(r, date.today())


@router.put("/{rid}")
def editar(rid: int, corpo: RestricaoIn, db: Session = Depends(get_db)):
    r = db.get(Restricao, rid)
    if r is None:
        raise HTTPException(404, "Restricao nao encontrada")
    _aplicar(r, corpo)
    db.commit()
    return _saida(r, date.today())


@router.delete("/{rid}", status_code=204)
def excluir(rid: int, db: Session = Depends(get_db)):
    r = db.get(Restricao, rid)
    if r is None:
        raise HTTPException(404, "Restricao nao encontrada")
    db.delete(r)
    db.commit()


# ---------------------------------------------------------------------- importacao


def _ler_planilha(nome: str, conteudo: bytes) -> list[list]:
    if nome.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        return [list(linha) for linha in wb.active.iter_rows(values_only=True)]
    for codificacao in ("utf-8-sig", "latin-1"):
        try:
            texto = conteudo.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    amostra = texto[:4096]
    delimitador = max([";", ",", "\t"], key=amostra.count)
    return [list(l) for l in csv.reader(io.StringIO(texto), delimiter=delimitador)]


@router.post("/importar/previa")
async def previa_importacao(arquivo: UploadFile = File(...)):
    """Le a planilha e sugere o mapeamento de colunas. Nao grava nada."""
    nome = arquivo.filename or ""
    if not nome.lower().endswith((".xlsx", ".xlsm", ".csv", ".txt")):
        raise HTTPException(400, "Envie um arquivo .xlsx ou .csv exportado do Prevision.")
    bruto = await arquivo.read()
    if len(bruto) > 10 * 1024 * 1024:
        raise HTTPException(400, "Arquivo acima de 10 MB.")
    try:
        linhas = [l for l in _ler_planilha(nome, bruto) if any(c not in (None, "") for c in l)]
    except Exception as e:  # noqa: BLE001
        raise HTTPException(400, f"Nao foi possivel ler a planilha: {type(e).__name__}") from e
    if len(linhas) < 2:
        raise HTTPException(400, "A planilha precisa de um cabecalho e pelo menos uma linha.")

    colunas = [str(c).strip() if c not in (None, "") else f"Coluna {i + 1}" for i, c in enumerate(linhas[0])]
    registros = []
    for linha in linhas[1:5001]:
        registros.append({
            col: (v.date().isoformat() if isinstance(v, datetime) else v.isoformat() if isinstance(v, date)
                  else "" if v is None else str(v).strip())
            for col, v in zip(colunas, linha)
        })

    usadas: set[str] = set()
    sugestao: dict[str, str | None] = {}
    for campo in CAMPOS:
        achada = None
        for col in colunas:
            if col not in usadas and any(s in _sem_acento(col) for s in SINONIMOS[campo]):
                achada = col
                break
        sugestao[campo] = achada
        if achada:
            usadas.add(achada)
    return {"arquivo": nome, "colunas": colunas, "total": len(registros), "amostra": registros[:6],
            "linhas": registros, "sugestao": sugestao}


class LoteIn(BaseModel):
    obra_id: int
    itens: list[dict]


@router.post("/importar/lote", status_code=201)
def importar_lote(corpo: LoteIn, db: Session = Depends(get_db)):
    if db.get(ObraCache, corpo.obra_id) is None:
        raise HTTPException(404, "Obra nao encontrada")
    criadas = ignoradas = 0
    for item in corpo.itens:
        descricao = str(item.get("descricao") or "").strip()
        if not descricao:
            ignoradas += 1
            continue
        status = _status(item.get("status"))
        r = Restricao(
            obra_id=corpo.obra_id, descricao=descricao,
            categoria=str(item.get("categoria") or "").strip() or None,
            responsavel=str(item.get("responsavel") or "").strip() or None,
            prazo=_data(item.get("prazo")), status=status,
            criada_em=_data(item.get("criada_em")) or date.today(),
            concluida_em=_data(item.get("concluida_em")) if status == "Concluida" else None,  # sem data na planilha: fica vazio, nao inventa
            origem="importacao", observacao=str(item.get("observacao") or "").strip() or None,
        )
        db.add(r)
        criadas += 1
    db.commit()
    return {"criadas": criadas, "ignoradas": ignoradas}
